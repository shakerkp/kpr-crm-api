import tkinter as tk
from psycopg2.extras import RealDictCursor
from tkinter import ttk, messagebox, filedialog
from db_manager import (
    get_db_connection, log_activity, get_customer_report_data,
    get_material_usage_report, get_material_transaction_details,
    get_material_usage_by_job, get_staff_usage_summary,
    get_all_users
)
import datetime
import csv
import openpyxl 

class ReportsModule:
    def __init__(self, parent_frame, current_user_info):
        self.parent_frame = parent_frame
        self.current_user_id = current_user_info.get('id')
        self.current_username = current_user_info.get('username')
        self.create_reports_ui()

    def create_reports_ui(self):
        # Clear existing widgets
        for widget in self.parent_frame.winfo_children():
            widget.destroy()

        # Title
        tk.Label(self.parent_frame, text="Reports & Analytics", font=("Arial", 20, "bold"), bg="#FFFFFF", fg="#333333").pack(pady=15)

        # Notebook for different report types
        self.report_notebook = ttk.Notebook(self.parent_frame)
        self.report_notebook.pack(pady=10, fill=tk.BOTH, expand=True, padx=10)

        # --- Tab 1: Financial Reports (Profit & Loss, Payments) ---
        financial_frame = ttk.Frame(self.report_notebook)
        self.report_notebook.add(financial_frame, text="Financial Reports")

        tk.Label(financial_frame, text="Financial Overview", font=("Arial", 16, "bold"), bg="#FFFFFF").pack(pady=10)

        date_range_frame = tk.Frame(financial_frame, bg="#FFFFFF")
        date_range_frame.pack(pady=5)

        tk.Label(date_range_frame, text="From:", bg="#FFFFFF").pack(side=tk.LEFT, padx=5)
        self.start_date_entry = ttk.Entry(date_range_frame, width=15)
        self.start_date_entry.pack(side=tk.LEFT, padx=5)
        self.start_date_entry.insert(0, (datetime.date.today() - datetime.timedelta(days=30)).strftime("%Y-%m-%d"))

        tk.Label(date_range_frame, text="To:", bg="#FFFFFF").pack(side=tk.LEFT, padx=5)
        self.end_date_entry = ttk.Entry(date_range_frame, width=15)
        self.end_date_entry.pack(side=tk.LEFT, padx=5)
        self.end_date_entry.insert(0, datetime.date.today().strftime("%Y-%m-%d"))

        tk.Button(date_range_frame, text="Generate Financial Report", command=self.generate_financial_report, bg="#007ACC", fg="white").pack(side=tk.LEFT, padx=10)

        self.financial_summary_label = tk.Label(financial_frame, text="", font=("Arial", 12), bg="#FFFFFF", justify=tk.LEFT)
        self.financial_summary_label.pack(pady=10, fill=tk.X)

        tk.Label(financial_frame, text="All Payments in Range", font=("Arial", 14, "bold"), bg="#FFFFFF").pack(pady=5, anchor="w")
        payment_cols = ("ID", "Customer", "Job ID", "Amount", "Date", "Mode")
        self.financial_payments_tree = ttk.Treeview(financial_frame, columns=payment_cols, show="headings")
        self.financial_payments_tree.pack(fill=tk.BOTH, expand=True)

        for col in payment_cols:
            self.financial_payments_tree.heading(col, text=col)
            self.financial_payments_tree.column(col, width=100)
        self.financial_payments_tree.column("ID", width=40)
        self.financial_payments_tree.column("Job ID", width=60)
        self.financial_payments_tree.column("Amount", width=80)
        
        payment_scrollbar = ttk.Scrollbar(financial_frame, orient="vertical", command=self.financial_payments_tree.yview)
        self.financial_payments_tree.configure(yscrollcommand=payment_scrollbar.set)
        payment_scrollbar.pack(side="right", fill="y")

        export_financial_frame = tk.Frame(financial_frame, bg="#FFFFFF")
        export_financial_frame.pack(pady=10)
        tk.Button(export_financial_frame, text="Export Payments to CSV", command=lambda: self.export_data_to_csv("payments"), bg="#28A745", fg="white").pack(side=tk.LEFT, padx=5)
        tk.Button(export_financial_frame, text="Export Payments to Excel", command=lambda: self.export_data_to_excel("payments"), bg="#28A745", fg="white").pack(side=tk.LEFT, padx=5)


        # --- Tab 2: Job Reports ---
        job_frame = ttk.Frame(self.report_notebook)
        self.report_notebook.add(job_frame, text="Job Reports")

        tk.Label(job_frame, text="Job Summary & Status", font=("Arial", 16, "bold"), bg="#FFFFFF").pack(pady=10)

        job_status_frame = tk.Frame(job_frame, bg="#FFFFFF")
        job_status_frame.pack(pady=5)

        tk.Label(job_status_frame, text="Filter by Status:", bg="#FFFFFF").pack(side=tk.LEFT, padx=5)
        self.job_status_var = tk.StringVar(job_status_frame)
        job_statuses = ["All", "Pending", "In Progress", "Completed", "Delivered"]
        self.job_status_dropdown = ttk.Combobox(job_status_frame, textvariable=self.job_status_var, values=job_statuses, state="readonly")
        self.job_status_dropdown.pack(side=tk.LEFT, padx=5)
        self.job_status_var.set("All")
        self.job_status_dropdown.bind("<<ComboboxSelected>>", self.load_jobs_report)

        tk.Button(job_status_frame, text="Refresh Job Report", command=self.load_jobs_report, bg="#007ACC", fg="white").pack(side=tk.LEFT, padx=10)

        self.job_summary_label = tk.Label(job_frame, text="", font=("Arial", 12), bg="#FFFFFF", justify=tk.LEFT)
        self.job_summary_label.pack(pady=10, fill=tk.X)

        tk.Label(job_frame, text="All Jobs", font=("Arial", 14, "bold"), bg="#FFFFFF").pack(pady=5, anchor="w")
        job_cols = ("ID", "Customer", "Type", "Status", "Price", "Start Date", "Delivery Date")
        self.jobs_tree = ttk.Treeview(job_frame, columns=job_cols, show="headings")
        self.jobs_tree.pack(fill=tk.BOTH, expand=True)

        for col in job_cols:
            self.jobs_tree.heading(col, text=col)
            self.jobs_tree.column(col, width=100)
        self.jobs_tree.column("ID", width=40)
        self.jobs_tree.column("Price", width=80)
        self.jobs_tree.column("Start Date", width=120)
        self.jobs_tree.column("Delivery Date", width=120)

        job_scrollbar = ttk.Scrollbar(job_frame, orient="vertical", command=self.jobs_tree.yview)
        self.jobs_tree.configure(yscrollcommand=job_scrollbar.set)
        job_scrollbar.pack(side="right", fill="y")

        export_job_frame = tk.Frame(job_frame, bg="#FFFFFF")
        export_job_frame.pack(pady=10)
        tk.Button(export_job_frame, text="Export Jobs to CSV", command=lambda: self.export_data_to_csv("jobs"), bg="#28A745", fg="white").pack(side=tk.LEFT, padx=5)
        tk.Button(export_job_frame, text="Export Jobs to Excel", command=lambda: self.export_data_to_excel("jobs"), bg="#28A745", fg="white").pack(side=tk.LEFT, padx=5)


        # --- Tab 3: Customer Reports ---
        customer_frame = ttk.Frame(self.report_notebook)
        self.report_notebook.add(customer_frame, text="Customer Reports")

        tk.Label(customer_frame, text="Customer Overview", font=("Arial", 16, "bold"), bg="#FFFFFF").pack(pady=10)

        customer_date_range_frame = tk.Frame(customer_frame, bg="#FFFFFF")
        customer_date_range_frame.pack(pady=5)

        tk.Label(customer_date_range_frame, text="From:", bg="#FFFFFF").pack(side=tk.LEFT, padx=5)
        self.customer_start_date_entry = ttk.Entry(customer_date_range_frame, width=15)
        self.customer_start_date_entry.pack(side=tk.LEFT, padx=5)
        self.customer_start_date_entry.insert(0, (datetime.date.today() - datetime.timedelta(days=365)).strftime("%Y-%m-%d"))

        tk.Label(customer_date_range_frame, text="To:", bg="#FFFFFF").pack(side=tk.LEFT, padx=5)
        self.customer_end_date_entry = ttk.Entry(customer_date_range_frame, width=15)
        self.customer_end_date_entry.pack(side=tk.LEFT, padx=5)
        self.customer_end_date_entry.insert(0, datetime.date.today().strftime("%Y-%m-%d"))

        customer_search_frame = tk.Frame(customer_frame, bg="#FFFFFF")
        customer_search_frame.pack(pady=5)
        tk.Label(customer_search_frame, text="Search Customer (Name/Mobile/Email):", bg="#FFFFFF").pack(side=tk.LEFT, padx=5)
        self.customer_search_var = tk.StringVar()
        self.customer_search_entry = ttk.Entry(customer_search_frame, textvariable=self.customer_search_var, width=40)
        self.customer_search_entry.pack(side=tk.LEFT, padx=5)

        tk.Button(customer_search_frame, text="Generate Customer Report", command=self.generate_customer_report, bg="#007ACC", fg="white").pack(side=tk.LEFT, padx=10)

        self.customer_summary_label = tk.Label(customer_frame, text="", font=("Arial", 12), bg="#FFFFFF", justify=tk.LEFT)
        self.customer_summary_label.pack(pady=10, fill=tk.X)

        tk.Label(customer_frame, text="All Customers", font=("Arial", 14, "bold"), bg="#FFFFFF").pack(pady=5, anchor="w")
        customer_cols = ("ID", "Name", "Mobile", "Email", "Total Jobs", "Total Billed", "Total Advance", "Balance Due", "Total Paid")
        self.customers_tree = ttk.Treeview(customer_frame, columns=customer_cols, show="headings")
        self.customers_tree.pack(fill=tk.BOTH, expand=True)

        for col in customer_cols:
            self.customers_tree.heading(col, text=col)
            self.customers_tree.column(col, width=100)
        self.customers_tree.column("ID", width=40)
        self.customers_tree.column("Mobile", width=100)
        self.customers_tree.column("Email", width=120)
        self.customers_tree.column("Total Jobs", width=80)
        
        customer_scrollbar = ttk.Scrollbar(customer_frame, orient="vertical", command=self.customers_tree.yview)
        self.customers_tree.configure(yscrollcommand=customer_scrollbar.set)
        customer_scrollbar.pack(side="right", fill="y")

        export_customer_frame = tk.Frame(customer_frame, bg="#FFFFFF")
        export_customer_frame.pack(pady=10)
        tk.Button(export_customer_frame, text="Export Customers to CSV", command=lambda: self.export_data_to_csv("customers"), bg="#28A745", fg="white").pack(side=tk.LEFT, padx=5)
        tk.Button(export_customer_frame, text="Export Customers to Excel", command=lambda: self.export_data_to_excel("customers"), bg="#28A745", fg="white").pack(side=tk.LEFT, padx=5)

        # --- Tab 4: Material Usage & Wastage Report ---
        material_frame = ttk.Frame(self.report_notebook)
        self.report_notebook.add(material_frame, text="📦 Material Usage")

        tk.Label(material_frame, text="Material Usage & Wastage Reports", font=("Arial", 16, "bold"),
                 bg="#FFFFFF").pack(pady=8)

        # ── Quick Date Preset Buttons ────────────────────────────────────────
        preset_outer = tk.Frame(material_frame, bg="#FFFFFF")
        preset_outer.pack(fill=tk.X, padx=10)

        tk.Label(preset_outer, text="Quick Range:", bg="#FFFFFF", font=("Arial", 9, "bold")).pack(side=tk.LEFT, padx=(0,5))
        for label, delta_days in [("Today", 0), ("Last 7 Days", 6), ("This Month", None), ("Last 30 Days", 29), ("Custom", -1)]:
            tk.Button(preset_outer, text=label, font=("Arial", 9),
                      bg="#E3F2FD", fg="#0D47A1", relief="ridge", padx=6,
                      command=lambda d=delta_days, lbl=label: self._apply_mat_preset(d, lbl)
                      ).pack(side=tk.LEFT, padx=3, pady=4)

        # ── Date + Staff filter bar ──────────────────────────────────────────
        mat_filter_frame = ttk.LabelFrame(material_frame, text="Filters", padding="8")
        mat_filter_frame.pack(fill=tk.X, padx=10, pady=4)

        ttk.Label(mat_filter_frame, text="Start Date:").grid(row=0, column=0, padx=5, sticky="w")
        self.mat_start_date = ttk.Entry(mat_filter_frame, width=13)
        self.mat_start_date.insert(0, (datetime.datetime.now() - datetime.timedelta(days=30)).strftime('%Y-%m-%d'))
        self.mat_start_date.grid(row=0, column=1, padx=5)

        ttk.Label(mat_filter_frame, text="End Date:").grid(row=0, column=2, padx=5, sticky="w")
        self.mat_end_date = ttk.Entry(mat_filter_frame, width=13)
        self.mat_end_date.insert(0, datetime.datetime.now().strftime('%Y-%m-%d'))
        self.mat_end_date.grid(row=0, column=3, padx=5)

        ttk.Label(mat_filter_frame, text="Staff:").grid(row=0, column=4, padx=5, sticky="w")
        self.mat_staff_var = tk.StringVar(value="All Staff")
        self.mat_staff_cb = ttk.Combobox(mat_filter_frame, textvariable=self.mat_staff_var,
                                         state="readonly", width=18)
        self._load_staff_filter_options()
        self.mat_staff_cb.grid(row=0, column=5, padx=5)

        ttk.Label(mat_filter_frame, text="Type:").grid(row=0, column=6, padx=5, sticky="w")
        self.mat_type_var = tk.StringVar(value="All")
        ttk.Combobox(mat_filter_frame, textvariable=self.mat_type_var,
                     values=["All", "Usage", "Wastage", "Stock In"],
                     state="readonly", width=10).grid(row=0, column=7, padx=5)

        ttk.Button(mat_filter_frame, text="🔍 Generate", command=self._generate_all_mat_reports).grid(row=0, column=8, padx=10)

        # ── Sub-tabs for the 4 different views ──────────────────────────────
        self.mat_notebook = ttk.Notebook(material_frame)
        self.mat_notebook.pack(fill=tk.BOTH, expand=True, padx=10, pady=6)

        # ── Sub-tab 1: Item Summary ──────────────────────────────────────────
        tab_summary = ttk.Frame(self.mat_notebook)
        self.mat_notebook.add(tab_summary, text="📊 Item Summary")

        sum_hdr = tk.Frame(tab_summary, bg="#FFFFFF")
        sum_hdr.pack(fill=tk.X)
        self.mat_summary_label = tk.Label(sum_hdr, text="", font=("Arial", 10), bg="#FFFFFF", fg="#333", justify=tk.LEFT)
        self.mat_summary_label.pack(side=tk.LEFT, padx=10, pady=4)
        ttk.Button(sum_hdr, text="📥 Export Excel", command=self.export_material_report).pack(side=tk.RIGHT, padx=5, pady=4)

        sum_tree_f = ttk.Frame(tab_summary)
        sum_tree_f.pack(fill=tk.BOTH, expand=True)
        mat_cols = ("SKU", "Item Name", "Type", "Size", "UoM", "Stock In", "Total Usage", "Total Wastage", "Net Consumed")
        self.mat_tree = ttk.Treeview(sum_tree_f, columns=mat_cols, show="headings")
        for col in mat_cols:
            self.mat_tree.heading(col, text=col)
            w = 150 if col == "Item Name" else 90
            self.mat_tree.column(col, width=w, anchor=tk.CENTER if col in ("SKU","UoM") else tk.W)
        self.mat_tree.tag_configure('high_wastage', background='#FFE0E0')
        self.mat_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        ttk.Scrollbar(sum_tree_f, orient="vertical", command=self.mat_tree.yview).pack(side=tk.RIGHT, fill=tk.Y)
        self.mat_tree.configure(yscrollcommand=lambda *a: None)

        # ── Sub-tab 2: Transaction Log (who entered what) ────────────────────
        tab_txn = ttk.Frame(self.mat_notebook)
        self.mat_notebook.add(tab_txn, text="📋 Transaction Log")

        txn_hdr = tk.Frame(tab_txn, bg="#FFFFFF")
        txn_hdr.pack(fill=tk.X)
        self.txn_count_label = tk.Label(txn_hdr, text="", font=("Arial", 10), bg="#FFFFFF", fg="#333")
        self.txn_count_label.pack(side=tk.LEFT, padx=10, pady=4)
        ttk.Button(txn_hdr, text="📥 Export Excel", command=self.export_transaction_log).pack(side=tk.RIGHT, padx=5, pady=4)

        txn_tree_f = ttk.Frame(tab_txn)
        txn_tree_f.pack(fill=tk.BOTH, expand=True)
        txn_cols = ("Date/Time", "Entered By", "Role", "Job#", "Customer", "Material", "Type", "Qty", "UoM", "Remarks")
        self.txn_tree = ttk.Treeview(txn_tree_f, columns=txn_cols, show="headings")
        col_widths = {"Date/Time": 130, "Entered By": 110, "Role": 70, "Job#": 55,
                      "Customer": 110, "Material": 140, "Type": 70, "Qty": 60, "UoM": 50, "Remarks": 150}
        for col in txn_cols:
            self.txn_tree.heading(col, text=col)
            self.txn_tree.column(col, width=col_widths.get(col, 90))
        self.txn_tree.tag_configure('wastage', foreground='#C62828')
        self.txn_tree.tag_configure('stock_in', foreground='#1B5E20')
        self.txn_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        txn_scroll = ttk.Scrollbar(txn_tree_f, orient="vertical", command=self.txn_tree.yview)
        txn_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.txn_tree.configure(yscrollcommand=txn_scroll.set)

        # ── Sub-tab 3: Per-Job Usage ─────────────────────────────────────────
        tab_job = ttk.Frame(self.mat_notebook)
        self.mat_notebook.add(tab_job, text="🔧 Per-Job Usage")

        job_hdr = tk.Frame(tab_job, bg="#FFFFFF")
        job_hdr.pack(fill=tk.X)
        self.job_mat_label = tk.Label(job_hdr, text="", font=("Arial", 10), bg="#FFFFFF", fg="#333")
        self.job_mat_label.pack(side=tk.LEFT, padx=10, pady=4)
        ttk.Button(job_hdr, text="📥 Export Excel", command=self.export_job_usage_report).pack(side=tk.RIGHT, padx=5, pady=4)

        job_tree_f = ttk.Frame(tab_job)
        job_tree_f.pack(fill=tk.BOTH, expand=True)
        job_cols = ("Job#", "Job Type", "Customer", "Description", "Material", "SKU", "UoM", "Usage", "Wastage")
        self.job_mat_tree = ttk.Treeview(job_tree_f, columns=job_cols, show="headings")
        job_col_w = {"Job#": 55, "Job Type": 110, "Customer": 120, "Description": 160,
                     "Material": 150, "SKU": 80, "UoM": 55, "Usage": 70, "Wastage": 70}
        for col in job_cols:
            self.job_mat_tree.heading(col, text=col)
            self.job_mat_tree.column(col, width=job_col_w.get(col, 90))
        self.job_mat_tree.tag_configure('has_wastage', background='#FFF3E0')
        self.job_mat_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        job_scroll = ttk.Scrollbar(job_tree_f, orient="vertical", command=self.job_mat_tree.yview)
        job_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.job_mat_tree.configure(yscrollcommand=job_scroll.set)

        # ── Sub-tab 4: Staff Accountability ─────────────────────────────────
        tab_staff = ttk.Frame(self.mat_notebook)
        self.mat_notebook.add(tab_staff, text="👤 Staff Summary")

        staff_hdr = tk.Frame(tab_staff, bg="#FFFFFF")
        staff_hdr.pack(fill=tk.X)
        self.staff_mat_label = tk.Label(staff_hdr, text="", font=("Arial", 10), bg="#FFFFFF", fg="#333")
        self.staff_mat_label.pack(side=tk.LEFT, padx=10, pady=4)
        ttk.Button(staff_hdr, text="📥 Export Excel", command=self.export_staff_summary).pack(side=tk.RIGHT, padx=5, pady=4)

        staff_tree_f = ttk.Frame(tab_staff)
        staff_tree_f.pack(fill=tk.BOTH, expand=True)
        staff_cols = ("Staff Name", "Role", "Jobs Worked", "Usage Entries", "Total Usage Qty", "Wastage Entries", "Total Wastage Qty")
        self.staff_mat_tree = ttk.Treeview(staff_tree_f, columns=staff_cols, show="headings")
        for col in staff_cols:
            self.staff_mat_tree.heading(col, text=col)
            self.staff_mat_tree.column(col, width=130 if "Name" in col else 105, anchor=tk.CENTER if "Qty" in col or "Entries" in col else tk.W)
        self.staff_mat_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        staff_scroll = ttk.Scrollbar(staff_tree_f, orient="vertical", command=self.staff_mat_tree.yview)
        staff_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.staff_mat_tree.configure(yscrollcommand=staff_scroll.set)

        # Initial loads
        self.generate_financial_report()
        self.load_jobs_report()
        self.generate_customer_report()
        self._generate_all_mat_reports()

    def generate_financial_report(self):
        start_date_str = self.start_date_entry.get()
        end_date_str = self.end_date_entry.get()

        conn = get_db_connection()
        cursor = conn.cursor(cursor_factory=RealDictCursor)

        cursor.execute("""
            SELECT SUM(amount) FROM payments
            WHERE payment_date BETWEEN ? AND ?
        """, (start_date_str + " 00:00:00", end_date_str + " 23:59:59"))
        total_payments = cursor.fetchone()[0] if cursor.rowcount else 0  # TODO: add AS alias or 0.0

        cursor.execute("""
            SELECT SUM(initial_price) FROM jobs
            WHERE (status = 'Completed' OR status = 'Delivered')
            AND delivery_date BETWEEN ? AND ?
        """, (start_date_str + " 00:00:00", end_date_str + " 23:59:59"))
        total_job_value = cursor.fetchone()[0] if cursor.rowcount else 0  # TODO: add AS alias or 0.0

        for item in self.financial_payments_tree.get_children():
            self.financial_payments_tree.delete(item)

        cursor.execute("""
            SELECT p.id, c.name AS customer_name, p.job_id, p.amount, p.payment_date, p.payment_mode
            FROM payments p
            JOIN customers c ON p.customer_id = c.id
            WHERE p.payment_date BETWEEN ? AND ?
            ORDER BY p.payment_date DESC
        """, (start_date_str + " 00:00:00", end_date_str + " 23:59:59"))
        payments = cursor.fetchall()
        conn.close()

        for pay in payments:
            self.financial_payments_tree.insert("", "end", values=(
                pay['id'], pay['customer_name'], pay['job_id'] if pay['job_id'] else "N/A",
                f"₹{pay['amount']:.2f}", pay['payment_date'], pay['payment_mode']
            ))

        summary_text = f"Report for: {start_date_str} to {end_date_str}\n\n" \
                       f"Total Payments Received: ₹{total_payments:.2f}\n" \
                       f"Estimated Job Value (Completed/Delivered): ₹{total_job_value:.2f}"
        self.financial_summary_label.config(text=summary_text)

    def load_jobs_report(self, event=None):
        status_filter = self.job_status_var.get()
        conn = get_db_connection()
        cursor = conn.cursor(cursor_factory=RealDictCursor)

        cursor.execute("SELECT status, COUNT(*) FROM jobs GROUP BY status")
        status_counts = {r[0]: r[1] for r in cursor.fetchall()}
        
        for item in self.jobs_tree.get_children():
            self.jobs_tree.delete(item)

        query = """
            SELECT j.id, c.name AS customer_name, j.job_type, j.status, j.initial_price, j.start_date, j.delivery_date
            FROM jobs j
            JOIN customers c ON j.customer_id = c.id
        """
        params = []
        if status_filter != "All":
            query += " WHERE j.status = ?"
            params.append(status_filter)
        query += " ORDER BY j.start_date DESC"

        cursor.execute(query, params)
        jobs = cursor.fetchall()
        conn.close()

        for job in jobs:
            self.jobs_tree.insert("", "end", values=(
                job['id'], job['customer_name'], job['job_type'], job['status'],
                f"₹{job['initial_price']:.2f}" if job['initial_price'] else "N/A",
                job['start_date'], job['delivery_date'] if job['delivery_date'] else "Pending"
            ))
        
        summary_text = "Job Status Summary:\n"
        total_jobs = 0
        for status in ["Pending", "In Progress", "Completed", "Delivered"]:
            count = status_counts.get(status, 0)
            summary_text += f"  {status}: {count}\n"
            total_jobs += count
        summary_text += f"Total Jobs: {total_jobs}"
        self.job_summary_label.config(text=summary_text)

    def generate_customer_report(self):
        start_date_str = self.customer_start_date_entry.get()
        end_date_str = self.customer_end_date_entry.get()
        search_term = self.customer_search_var.get().strip()

        for item in self.customers_tree.get_children():
            self.customers_tree.delete(item)

        customers_data = get_customer_report_data(start_date_str, end_date_str, search_term)
        
        total_customers = len(customers_data)
        customers_with_jobs = sum(1 for c in customers_data if c['total_jobs'] > 0)
        customers_with_payments = sum(1 for c in customers_data if c['total_payments_received'] and c['total_payments_received'] > 0)

        for customer in customers_data:
            self.customers_tree.insert("", "end", values=(
                customer['id'],
                customer['name'],
                customer['mobile'],
                customer['email'] if customer['email'] else "N/A",
                customer['total_jobs'] if customer['total_jobs'] else 0,
                f"₹{customer['total_billed_amount']:.2f}" if customer['total_billed_amount'] else "₹0.00",
                f"₹{customer['total_advance_received']:.2f}" if customer['total_advance_received'] else "₹0.00",
                f"₹{customer['total_balance_due']:.2f}" if customer['total_balance_due'] else "₹0.00",
                f"₹{customer['total_payments_received']:.2f}" if customer['total_payments_received'] else "₹0.00"
            ))
        
        summary_text = f"Report for: {start_date_str} to {end_date_str}\n" \
                       f"Total Customers: {total_customers}\n" \
                       f"Customers with Jobs: {customers_with_jobs}\n" \
                       f"Customers with Payments: {customers_with_payments}"
        self.customer_summary_label.config(text=summary_text)

    # ─────────────────────────────────────────────────────────────────────────
    # MATERIAL USAGE REPORT HELPERS
    # ─────────────────────────────────────────────────────────────────────────

    def _load_staff_filter_options(self):
        """Populate the Staff filter combobox from users table."""
        try:
            all_users = get_all_users()
            names = ["All Staff"] + [
                f"{u.get('full_name') or u['username']} ({u['username']})"
                for u in all_users if u.get('role') in ('staff', 'sub_staff', 'admin', 'user')
            ]
            self._staff_options_map = {
                f"{u.get('full_name') or u['username']} ({u['username']})": u['id']
                for u in all_users
            }
            self.mat_staff_cb['values'] = names
        except Exception:
            self.mat_staff_cb['values'] = ["All Staff"]
            self._staff_options_map = {}

    def _apply_mat_preset(self, delta_days, label):
        today = datetime.date.today()
        if label == "Today":
            self.mat_start_date.delete(0, tk.END)
            self.mat_start_date.insert(0, today.strftime('%Y-%m-%d'))
            self.mat_end_date.delete(0, tk.END)
            self.mat_end_date.insert(0, today.strftime('%Y-%m-%d'))
        elif label == "This Month":
            start = today.replace(day=1)
            self.mat_start_date.delete(0, tk.END)
            self.mat_start_date.insert(0, start.strftime('%Y-%m-%d'))
            self.mat_end_date.delete(0, tk.END)
            self.mat_end_date.insert(0, today.strftime('%Y-%m-%d'))
        elif delta_days >= 0:
            start = today - datetime.timedelta(days=delta_days)
            self.mat_start_date.delete(0, tk.END)
            self.mat_start_date.insert(0, start.strftime('%Y-%m-%d'))
            self.mat_end_date.delete(0, tk.END)
            self.mat_end_date.insert(0, today.strftime('%Y-%m-%d'))
        # "Custom" — user edits manually, do nothing
        self._generate_all_mat_reports()

    def _get_mat_filters(self):
        """Return (start_str, end_str, user_id_filter, type_filter)."""
        start = self.mat_start_date.get().strip()
        end = self.mat_end_date.get().strip()
        staff_sel = self.mat_staff_var.get()
        user_id_filter = self._staff_options_map.get(staff_sel, None) if hasattr(self, '_staff_options_map') else None
        type_filter = self.mat_type_var.get()
        return start, end, user_id_filter, type_filter

    def _generate_all_mat_reports(self):
        """Regenerate all 4 material sub-tabs at once."""
        self.load_material_report()
        self._load_transaction_log()
        self._load_job_usage_report()
        self._load_staff_usage_report()

    def load_material_report(self):
        """Sub-tab 1: Item-level summary (Stock In / Usage / Wastage)."""
        for row in self.mat_tree.get_children():
            self.mat_tree.delete(row)

        start, end, _, _ = self._get_mat_filters()
        data = get_material_usage_report(start, end)

        total_usage = 0.0
        total_wastage = 0.0
        for item in data:
            u = item.get('total_usage') or 0
            w = item.get('total_wastage') or 0
            net = u + w
            total_usage += u
            total_wastage += w
            tag = ('high_wastage',) if w > 0 and u > 0 and (w / (u + w)) > 0.3 else ()
            self.mat_tree.insert("", tk.END, values=(
                item.get('sku') or "-",
                item['item_name'],
                item.get('item_type') or "-",
                item.get('item_size') or "-",
                item.get('uom') or "Pcs",
                item.get('total_inward') or 0,
                round(u, 3),
                round(w, 3),
                round(net, 3)
            ), tags=tag)

        self.mat_summary_label.config(
            text=f"Period: {start}  →  {end}   |   Items: {len(data)}   |   "
                 f"Total Usage: {round(total_usage,3)}   |   Total Wastage: {round(total_wastage,3)}"
        )

    def _load_transaction_log(self):
        """Sub-tab 2: Full per-transaction audit log."""
        for row in self.txn_tree.get_children():
            self.txn_tree.delete(row)

        start, end, user_id_filter, type_filter = self._get_mat_filters()
        data = get_material_transaction_details(start, end, user_id_filter, type_filter)

        for txn in data:
            tag = ()
            if txn['transaction_type'] == 'Wastage':
                tag = ('wastage',)
            elif txn['transaction_type'] == 'Stock In':
                tag = ('stock_in',)
            job_disp = f"#{txn['job_id']}" if txn.get('job_id') else "-"
            self.txn_tree.insert("", tk.END, values=(
                txn['timestamp'][:16],
                txn['entered_by'],
                txn.get('user_role', ''),
                job_disp,
                txn.get('customer_name', '-'),
                txn['material_name'],
                txn['transaction_type'],
                round(txn['quantity'], 3),
                txn.get('uom', 'Pcs'),
                txn.get('remarks', '')
            ), tags=tag)

        self.txn_count_label.config(text=f"Transactions: {len(data)}   |   Period: {start} → {end}")

    def _load_job_usage_report(self):
        """Sub-tab 3: Per-job material consumption."""
        for row in self.job_mat_tree.get_children():
            self.job_mat_tree.delete(row)

        start, end, _, _ = self._get_mat_filters()
        data = get_material_usage_by_job(start, end)

        total_jobs = len({r['job_id'] for r in data if r['job_id']})
        for row in data:
            w = row.get('total_wastage') or 0
            tag = ('has_wastage',) if w > 0 else ()
            job_disp = f"#{row['job_id']}" if row.get('job_id') else "No Job"
            self.job_mat_tree.insert("", tk.END, values=(
                job_disp,
                row.get('job_type', '-'),
                row.get('customer_name', '-'),
                (row.get('job_description') or '')[:40],
                row['material_name'],
                row.get('sku', '-'),
                row.get('uom', 'Pcs'),
                round(row.get('total_usage') or 0, 3),
                round(w, 3)
            ), tags=tag)

        self.job_mat_label.config(
            text=f"Period: {start} → {end}   |   Jobs with material entries: {total_jobs}   |   Rows: {len(data)}"
        )

    def _load_staff_usage_report(self):
        """Sub-tab 4: Staff accountability summary."""
        for row in self.staff_mat_tree.get_children():
            self.staff_mat_tree.delete(row)

        start, end, _, _ = self._get_mat_filters()
        data = get_staff_usage_summary(start, end)

        for row in data:
            self.staff_mat_tree.insert("", tk.END, values=(
                row.get('staff_name', 'Unknown'),
                row.get('role', ''),
                row.get('jobs_worked') or 0,
                row.get('usage_entries') or 0,
                round(row.get('total_usage_qty') or 0, 3),
                row.get('wastage_entries') or 0,
                round(row.get('total_wastage_qty') or 0, 3)
            ))

        self.staff_mat_label.config(
            text=f"Period: {start} → {end}   |   Staff with entries: {len(data)}"
        )

    # ── Export helpers ────────────────────────────────────────────────────────

    def _export_treeview_to_excel(self, tree, headers, sheet_title, filename_prefix):
        """Generic helper: exports any Treeview to an Excel file."""
        if openpyxl is None:
            messagebox.showerror("Error", "Please install 'openpyxl' for Excel export.")
            return
        rows = [tree.item(r, 'values') for r in tree.get_children()]
        if not rows:
            messagebox.showinfo("No Data", "No data to export.")
            return
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
            initialfile=f"{filename_prefix}_{datetime.datetime.now().strftime('%Y%m%d')}.xlsx"
        )
        if not file_path:
            return
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_title
        ws.append(headers)
        for row in rows:
            ws.append(list(row))
        wb.save(file_path)
        messagebox.showinfo("Export Success", f"Exported to:\n{file_path}")

    def export_material_report(self):
        """Export Item Summary tab."""
        start, end, _, _ = self._get_mat_filters()
        data = get_material_usage_report(start, end)
        if openpyxl is None:
            messagebox.showerror("Error", "Please install 'openpyxl'."); return
        if not data:
            messagebox.showinfo("No Data", "No material data found."); return
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", filetypes=[("Excel Files", "*.xlsx")],
            title="Save Material Summary Report"
        )
        if not file_path: return
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Material Summary"
        ws.append(["SKU", "Item Name", "Type", "Size", "UoM", "Stock In", "Total Usage", "Total Wastage", "Net Consumed"])
        for item in data:
            u = item.get('total_usage') or 0
            w = item.get('total_wastage') or 0
            ws.append([item.get('sku'), item['item_name'], item.get('item_type'), item.get('item_size'),
                       item.get('uom'), item.get('total_inward'), round(u,3), round(w,3), round(u+w,3)])
        wb.save(file_path)
        messagebox.showinfo("Export Success", f"Saved to:\n{file_path}")

    def export_transaction_log(self):
        self._export_treeview_to_excel(
            self.txn_tree,
            ["Date/Time", "Entered By", "Role", "Job#", "Customer", "Material", "Type", "Qty", "UoM", "Remarks"],
            "Transaction Log", "Transaction_Log"
        )

    def export_job_usage_report(self):
        self._export_treeview_to_excel(
            self.job_mat_tree,
            ["Job#", "Job Type", "Customer", "Description", "Material", "SKU", "UoM", "Usage", "Wastage"],
            "Per-Job Usage", "Job_Material_Usage"
        )

    def export_staff_summary(self):
        self._export_treeview_to_excel(
            self.staff_mat_tree,
            ["Staff Name", "Role", "Jobs Worked", "Usage Entries", "Total Usage Qty", "Wastage Entries", "Total Wastage Qty"],
            "Staff Summary", "Staff_Material_Summary"
        )

    def export_data_to_csv(self, report_type):
        file_path = filedialog.asksaveasfilename(defaultextension=".csv", filetypes=[("CSV files", "*.csv"), ("All files", "*.*")])
        if not file_path:
            return

        conn = get_db_connection()
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        data = []
        headers = []

        if report_type == "payments":
            headers = ["ID", "Customer Name", "Job ID", "Amount", "Payment Date", "Payment Mode", "Notes"]
            cursor.execute("SELECT p.id, c.name, p.job_id, p.amount, p.payment_date, p.payment_mode, p.notes FROM payments p JOIN customers c ON p.customer_id = c.id ORDER BY p.payment_date DESC")
            data = cursor.fetchall()
        elif report_type == "jobs":
            headers = ["ID", "Customer Name", "Job Type", "Description", "Initial Price", "Assigned Staff ID", "Status", "Start Date", "Delivery Date"]
            cursor.execute("SELECT j.id, c.name, j.job_type, j.description, j.initial_price, j.assigned_staff_id, j.status, j.start_date, j.delivery_date FROM jobs j JOIN customers c ON j.customer_id = c.id ORDER BY j.start_date DESC")
            data = cursor.fetchall()
        elif report_type == "customers":
            headers = ["ID", "Name", "Mobile", "Email", "Address", "Tags", "Total Jobs", "Total Billed Amount", "Total Advance Received", "Total Balance Due", "Total Payments Received"]
            data = get_customer_report_data(self.customer_start_date_entry.get(), self.customer_end_date_entry.get(), self.customer_search_var.get().strip())
        
        conn.close()
        
        try:
            with open(file_path, 'w', newline='', encoding='utf-8') as csvfile:
                csv_writer = csv.writer(csvfile)
                csv_writer.writerow(headers)
                for row in data:
                    csv_writer.writerow(list(row.values()) if isinstance(row, dict) else list(row))
            messagebox.showinfo("Export Success", f"{report_type.capitalize()} data exported to CSV successfully.")
        except Exception as e:
            messagebox.showerror("Export Error", f"Failed to export data: {e}")

    def export_data_to_excel(self, report_type):
        if openpyxl is None:
            messagebox.showerror("Error", "Please install 'openpyxl'.")
            return
            
        file_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")])
        if not file_path:
            return

        conn = get_db_connection()
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        data = []
        headers = []

        if report_type == "payments":
            headers = ["ID", "Customer Name", "Job ID", "Amount", "Payment Date", "Payment Mode", "Notes"]
            cursor.execute("SELECT p.id, c.name, p.job_id, p.amount, p.payment_date, p.payment_mode, p.notes FROM payments p JOIN customers c ON p.customer_id = c.id ORDER BY p.payment_date DESC")
            data = cursor.fetchall()
        elif report_type == "jobs":
            headers = ["ID", "Customer Name", "Job Type", "Description", "Initial Price", "Status", "Start Date", "Delivery Date"]
            cursor.execute("SELECT j.id, c.name, j.job_type, j.description, j.initial_price, j.status, j.start_date, j.delivery_date FROM jobs j JOIN customers c ON j.customer_id = c.id ORDER BY j.start_date DESC")
            data = cursor.fetchall()
        elif report_type == "customers":
            headers = ["ID", "Name", "Mobile", "Email", "Address", "Total Jobs", "Total Billed", "Total Advance", "Balance Due", "Total Paid"]
            data = get_customer_report_data(self.customer_start_date_entry.get(), self.customer_end_date_entry.get(), self.customer_search_var.get().strip())
            
        conn.close()

        try:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = f"{report_type.capitalize()} Report"
            sheet.append(headers)
            for row in data:
                sheet.append(list(row.values()) if isinstance(row, dict) else list(row))
            workbook.save(file_path)
            messagebox.showinfo("Export Success", f"{report_type.capitalize()} data exported to Excel.")
        except Exception as e:
            messagebox.showerror("Export Error", f"Failed to export data: {e}")